import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from app.storage.case_matrix_repository import CaseMatrixRepository


class CaseMatrixRepositoryTests(unittest.TestCase):
    def test_reads_case_results_from_scenario_header(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            matrix_path = Path(temporary_directory) / "matrix.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append([
                "ID",
                "Escenario datos de Prueba",
                "Resultado esperado",
                "Resultado obtenido",
            ])
            worksheet.append(["CP01", "Inicio de sesión", "Acceso permitido", "Ingresó al sistema"])
            worksheet.append(["CP02", "Recuperar contraseña", "Correo enviado", "Pendiente"])
            workbook.save(matrix_path)
            workbook.close()

            cases = CaseMatrixRepository.load_cases(matrix_path)
            self.assertEqual(len(cases), 2)
            self.assertEqual(cases[0].name, "Inicio de sesión")
            self.assertEqual(cases[0].expected_result, "Acceso permitido")
            self.assertEqual(cases[0].obtained_result, "Ingresó al sistema")
            self.assertEqual(cases[0].case_id, "CP01")
            self.assertEqual(cases[1].name, "Recuperar contraseña")
            self.assertNotEqual(cases[0].matrix_key, cases[1].matrix_key)

    def test_reports_when_no_case_name_column_is_found(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            matrix_path = Path(temporary_directory) / "matrix.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["ID", "Resultado esperado"])
            worksheet.append(["CP01", "Acceso permitido"])
            workbook.save(matrix_path)
            workbook.close()

            with self.assertRaisesRegex(ValueError, "No se encontró una columna"):
                CaseMatrixRepository.load_case_names(matrix_path)

    def test_expected_result_uses_operational_columns_that_have_values(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            matrix_path = Path(temporary_directory) / "matrix.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append([
                "Escenario datos de Prueba",
                "Resultados operativos",
                "Operativo contable",
            ])
            worksheet.append(["Caso completo", "Saldo correcto", "Débito y crédito cuadran"])
            worksheet.append(["Solo contable", None, "Asiento contabilizado"])
            worksheet.append(["Sin resultado", None, None])
            workbook.save(matrix_path)
            workbook.close()

            cases = CaseMatrixRepository.load_cases(matrix_path)

            self.assertEqual(
                cases[0].expected_result,
                "Saldo correcto\nDébito y crédito cuadran",
            )
            self.assertEqual(cases[1].expected_result, "Asiento contabilizado")
            self.assertEqual(cases[2].expected_result, "")


if __name__ == "__main__":
    unittest.main()