import importlib.util
import tempfile
import unittest
from pathlib import Path

from app.reporting.test_case_report import TestCaseReportGenerator
from app.ui.main_window import MainWindow


@unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx no está instalado")
class ReportTargetTests(unittest.TestCase):
    def test_report_document_names_lists_existing_word_files(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            (folder / "Regresion.docx").touch()
            (folder / "not-a-report.txt").touch()

            self.assertEqual(
                MainWindow._report_document_names(folder),
                ["Regresion.docx"],
            )

    def test_next_available_report_name_avoids_overwriting(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            (folder / "Regresion.docx").touch()
            (folder / "Regresion_nuevo.docx").touch()

            self.assertEqual(
                MainWindow._next_available_report_name(folder, "Regresion.docx"),
                "Regresion_nuevo_2.docx",
            )

    def test_existing_selected_word_requires_maintain_checkbox_to_append(self):
        self.assertTrue(
            MainWindow._requires_reuse_checkbox(
                folder_exists=True,
                report_exists=True,
                creating_new_folder=False,
                creating_new_document=False,
                maintain_target=False,
            )
        )
        self.assertFalse(
            MainWindow._requires_reuse_checkbox(
                folder_exists=True,
                report_exists=True,
                creating_new_folder=False,
                creating_new_document=False,
                maintain_target=True,
            )
        )
        self.assertTrue(
            MainWindow._requires_reuse_checkbox(
                folder_exists=True,
                report_exists=True,
                creating_new_folder=False,
                creating_new_document=True,
                maintain_target=False,
            )
        )

    def test_new_folder_and_word_can_be_created_for_first_run(self):
        self.assertFalse(
            MainWindow._requires_reuse_checkbox(
                folder_exists=False,
                report_exists=False,
                creating_new_folder=True,
                creating_new_document=True,
                maintain_target=False,
            )
        )

    def test_reusing_a_created_new_folder_requires_maintain_checkbox(self):
        self.assertTrue(
            MainWindow._requires_reuse_checkbox(
                folder_exists=True,
                report_exists=True,
                creating_new_folder=True,
                creating_new_document=True,
                maintain_target=False,
            )
        )

    def test_existing_folder_with_new_word_requires_maintain_checkbox(self):
        self.assertTrue(
            MainWindow._requires_reuse_checkbox(
                folder_exists=True,
                report_exists=False,
                creating_new_folder=False,
                creating_new_document=True,
                maintain_target=False,
            )
        )

    def test_case_name_is_in_table_without_heading_above_it(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            execution_folder = root / "execution"
            execution_folder.mkdir()
            report_folder = root / "reports"

            output_file = TestCaseReportGenerator(report_folder).generate_or_append(
                execution_folder,
                [],
                case_name="Compra de prueba",
                report_name="caso.docx",
            )
            document = Document(output_file)

            self.assertFalse(
                any("Compra de prueba" in paragraph.text for paragraph in document.paragraphs)
            )
            table_text = " ".join(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn("Compra de prueba", table_text)

    def test_custom_word_name_is_reused_for_appended_cases(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            execution_folder = root / "execution"
            report_folder = root / "reports" / "Sprint 1"
            execution_folder.mkdir()
            report_name = "Regresion UI.docx"
            generator = TestCaseReportGenerator(report_folder)

            first_report = generator.generate_or_append(
                execution_folder,
                [],
                case_name="Inicio de sesión",
                report_name=report_name,
            )
            second_report = generator.generate_or_append(
                execution_folder,
                [],
                case_name="Cierre de sesión",
                report_name=report_name,
            )

            self.assertEqual(first_report, report_folder / report_name)
            self.assertEqual(second_report, first_report)
            self.assertTrue(second_report.exists())
            document = Document(second_report)
            self.assertEqual(len(document.tables), 2)
            self.assertEqual(len(document.paragraphs), 1)
            self.assertEqual(document.paragraphs[0].text, "")
            case_ids = [
                table.rows[0].cells[0].text.split(" - ", maxsplit=1)[0]
                for table in document.tables
            ]
            self.assertEqual(case_ids, ["CP01", "CP02"])

    def test_matrix_case_can_append_new_id_or_overwrite_its_existing_table(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            execution_folder = root / "execution"
            execution_folder.mkdir()
            report_folder = root / "reports"
            generator = TestCaseReportGenerator(report_folder)
            matrix_key = "matrix.xlsx|Casos|2"

            generator.generate_or_append(
                execution_folder,
                [],
                case_name="Inicio de sesión",
                report_name="casos.docx",
                matrix_case_key=matrix_key,
                matrix_case_id="MAT-01",
                expected_result="Se inicia sesión correctamente",
                obtained_result="Se inicia sesión correctamente",
            )
            report_file = report_folder / "casos.docx"
            self.assertEqual(
                generator.find_matrix_case_id(report_file, matrix_key),
                "CP01",
            )

            generator.generate_or_append(
                execution_folder,
                [],
                case_name="Inicio de sesión",
                report_name="casos.docx",
                matrix_case_key=matrix_key,
                matrix_case_id="MAT-01",
                expected_result="Se inicia sesión correctamente",
                obtained_result="Falló el primer intento",
            )
            self.assertEqual(
                generator.find_matrix_case_id(report_file, matrix_key),
                "CP02",
            )

            generator.generate_or_append(
                execution_folder,
                [],
                case_name="Inicio de sesión actualizado",
                report_name="casos.docx",
                matrix_case_key=matrix_key,
                matrix_case_id="MAT-01",
                expected_result="Se inicia sesión correctamente, sin errores",
                obtained_result="Resultado corregido",
                overwrite_matrix_case=True,
            )

            document = Document(report_file)
            self.assertEqual(len(document.tables), 2)
            self.assertEqual(
                [table.rows[0].cells[0].text.split(" - ", maxsplit=1)[0]
                 for table in document.tables],
                ["CP01", "CP02"],
            )
            self.assertIn("MAT-01", document.tables[1].cell(0, 0).text + " ".join(
                cell.text for row in document.tables[1].rows for cell in row.cells
            ))
            self.assertIn("Resultado corregido", " ".join(
                cell.text for row in document.tables[1].rows for cell in row.cells
            ))
            self.assertIn("Se inicia sesión correctamente, sin errores", " ".join(
                cell.text for row in document.tables[1].rows for cell in row.cells
            ))
            self.assertNotIn("Falló el primer intento", " ".join(
                cell.text for row in document.tables[1].rows for cell in row.cells
            ))

    def test_missing_matrix_expected_result_uses_default_text(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            execution_folder = root / "execution"
            execution_folder.mkdir()
            report_folder = root / "reports"

            report_file = TestCaseReportGenerator(report_folder).generate_or_append(
                execution_folder,
                [],
                case_name="Caso sin resultado esperado",
                report_name="casos.docx",
                matrix_case_key="matrix.xlsx|Casos|2",
            )
            document = Document(report_file)
            table_text = " ".join(
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )

            self.assertIn("No se tiene en la matriz.", table_text)


if __name__ == "__main__":
    unittest.main()