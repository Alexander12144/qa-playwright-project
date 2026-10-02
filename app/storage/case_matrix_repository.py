from pathlib import Path
import re
import unicodedata
from dataclasses import dataclass

from openpyxl import load_workbook


@dataclass(frozen=True)
class MatrixCase:
    name: str
    expected_result: str = ""
    obtained_result: str = ""
    case_id: str = ""
    matrix_key: str = ""


class CaseMatrixRepository:
    """Lee los nombres de caso desde una matriz Excel."""

    HEADER_SCAN_ROWS = 30
    EXACT_HEADERS = {
        "case name": 0,
        "test case name": 0,
        "nombre caso de prueba": 0,
        "nombre del caso de prueba": 0,
        "nombre del caso": 1,
        "nombre caso": 1,
        "caso de prueba": 2,
        "test case": 2,
        "escenario de prueba": 3,
        "escenario datos de prueba": 0,
        "scenario": 3,
    }
    FIELD_HEADERS = {
        "expected_result": {
            "resultado esperado",
            "resultado esperado del caso",
            "resultados operativos",
            "operativo contable",
            "expected result",
            "expected outcome",
        },
        "obtained_result": {
            "resultado obtenido",
            "resultado real",
            "resultado actual",
            "actual result",
            "obtained result",
        },
        "case_id": {
            "id",
            "id caso",
            "id del caso",
            "identificador",
            "case id",
            "test case id",
        },
    }

    @classmethod
    def load_cases(cls, filepath: Path | str) -> list[MatrixCase]:
        filepath = Path(filepath)
        workbook = load_workbook(filepath, read_only=True, data_only=True)
        try:
            candidates = []
            for sheet_index, worksheet in enumerate(workbook.worksheets):
                for row_index, row in enumerate(
                    worksheet.iter_rows(
                        min_row=1,
                        max_row=cls.HEADER_SCAN_ROWS,
                        values_only=True,
                    ),
                    start=1,
                ):
                    for column_index, value in enumerate(row):
                        score = cls._header_score(value)
                        if score is not None:
                            candidates.append(
                                (score, sheet_index, row_index, column_index, worksheet)
                            )

            if not candidates:
                raise ValueError(
                    "No se encontró una columna de casos. Use un encabezado "
                    "como 'Nombre del caso de prueba' o 'Caso de prueba'."
                )

            _, sheet_index, header_row, name_column, worksheet = min(
                candidates,
                key=lambda candidate: candidate[:4],
            )
            header_columns: dict[str, list[int]] = {}
            header_values = next(
                worksheet.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    values_only=True,
                )
            )
            for column_index, value in enumerate(header_values):
                field_name = cls._field_name(value)
                if field_name:
                    header_columns.setdefault(field_name, []).append(column_index)

            matrix_cases = []
            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=header_row + 1,
                    values_only=True,
                ),
                start=header_row + 1,
            ):
                name_value = row[name_column] if name_column < len(row) else None
                if name_value is None or not str(name_value).strip():
                    continue

                def field_value(field_name: str) -> str:
                    values = []
                    for column_index in header_columns.get(field_name, []):
                        if column_index >= len(row) or row[column_index] is None:
                            continue
                        value = str(row[column_index]).strip()
                        if value:
                            values.append(value)
                    return "\n".join(values)

                matrix_cases.append(
                    MatrixCase(
                        name=str(name_value).strip(),
                        expected_result=field_value("expected_result"),
                        obtained_result=field_value("obtained_result"),
                        case_id=field_value("case_id"),
                        matrix_key=(
                            f"{filepath.name}|{worksheet.title}|{row_number}"
                        ),
                    )
                )

            if not matrix_cases:
                raise ValueError("La columna de casos no contiene nombres.")
            return matrix_cases
        finally:
            workbook.close()

    @classmethod
    def load_case_names(cls, filepath: Path | str) -> list[str]:
        return [case.name for case in cls.load_cases(filepath)]

    @classmethod
    def _field_name(cls, value) -> str | None:
        normalized = cls._normalize_header(value)
        for field_name, headers in cls.FIELD_HEADERS.items():
            if normalized in headers:
                return field_name
        return None

    @staticmethod
    def _normalize_header(value) -> str:
        if value is None:
            return ""
        normalized = unicodedata.normalize("NFKD", str(value))
        normalized = "".join(
            character
            for character in normalized
            if not unicodedata.combining(character)
        )
        return re.sub(r"[^a-zA-Z0-9]+", " ", normalized).strip().lower()

    @classmethod
    def _header_score(cls, value) -> int | None:
        normalized = cls._normalize_header(value)
        if not normalized:
            return None
        if normalized in cls.EXACT_HEADERS:
            return cls.EXACT_HEADERS[normalized]
        if "nombre" in normalized and "caso" in normalized:
            return 1
        if "case" in normalized and ("name" in normalized or "test" in normalized):
            return 1
        return None