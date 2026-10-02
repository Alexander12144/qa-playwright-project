from datetime import datetime
from pathlib import Path
import re
from urllib.parse import urlparse

class TestCaseReportGenerator:
    """Genera un caso de prueba Word a partir de acciones QA."""

    REPORT_FOLDER = Path("data") / "test_cases"

    def __init__(self, report_folder: Path | None = None) -> None:
        self.report_folder = report_folder or self.REPORT_FOLDER

    def generate(
        self,
        execution_folder: Path,
        qa_actions: list[dict],
        start_url: str | None = None,
        case_name: str | None = None,
        report_name: str = "CasosDePrueba.docx",
        matrix_case_key: str | None = None,
        matrix_case_id: str | None = None,
        expected_result: str | None = None,
        obtained_result: str | None = None,
        overwrite_matrix_case: bool = False,
    ) -> Path:
        return self.generate_or_append(
            execution_folder,
            qa_actions,
            start_url=start_url,
            case_name=case_name,
            report_name=report_name,
            matrix_case_key=matrix_case_key,
            matrix_case_id=matrix_case_id,
            expected_result=expected_result,
            obtained_result=obtained_result,
            overwrite_matrix_case=overwrite_matrix_case,
        )

    def generate_or_append(
        self,
        execution_folder: Path,
        qa_actions: list[dict],
        start_url: str | None = None,
        case_name: str | None = None,
        report_name: str = "CasosDePrueba.docx",
        matrix_case_key: str | None = None,
        matrix_case_id: str | None = None,
        expected_result: str | None = None,
        obtained_result: str | None = None,
        overwrite_matrix_case: bool = False,
    ) -> Path:
        try:
            from docx import Document
            from docx.enum.table import WD_TABLE_ALIGNMENT
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.shared import Inches, Pt, RGBColor
        except ModuleNotFoundError as error:
            raise RuntimeError(
                "Falta python-docx. Ejecuta: "
                "python -m pip install -r requirements.txt"
            ) from error

        self._Document = Document
        self._WD_TABLE_ALIGNMENT = WD_TABLE_ALIGNMENT
        self._OxmlElement = OxmlElement
        self._qn = qn
        self._Inches = Inches
        self._Pt = Pt
        self._RGBColor = RGBColor

        self.report_folder.mkdir(parents=True, exist_ok=True)
        output_file = self.report_folder / report_name
        existing_case_table = None

        if output_file.exists():
            document = self._Document(output_file)
            if matrix_case_key:
                existing_case_table = self._find_matrix_case_table(
                    document,
                    matrix_case_key,
                )
            if overwrite_matrix_case and existing_case_table is not None:
                case_id = self._table_case_id(existing_case_table)
            else:
                case_id = self._build_case_id(document)
        else:
            document = self._Document()
            self._configure_document(document)
            case_id = self._build_case_id(document)

        new_table, separator = self._add_case_table(
            document,
            case_id,
            execution_folder,
            qa_actions,
            start_url,
            case_name=case_name,
            matrix_case_key=matrix_case_key,
            matrix_case_id=matrix_case_id,
            expected_result=expected_result,
            obtained_result=obtained_result,
        )

        if overwrite_matrix_case and existing_case_table is not None:
            old_table = existing_case_table._tbl
            if separator is not None:
                old_table.addprevious(separator._p)
            old_table.addprevious(new_table._tbl)
            old_table.getparent().remove(old_table)

        document.save(output_file)
        return output_file

    @classmethod
    def find_matrix_case_id(cls, report_file: Path, matrix_case_key: str) -> str | None:
        if not report_file.exists():
            return None
        from docx import Document

        document = Document(report_file)
        matches = [
            table
            for table in document.tables
            if cls._table_matrix_key(table) == matrix_case_key
        ]
        return cls._table_case_id(matches[-1]) if matches else None

    def _find_matrix_case_table(self, document, matrix_case_key: str):
        matches = [
            table
            for table in document.tables
            if self._table_matrix_key(table) == matrix_case_key
        ]
        return matches[-1] if matches else None

    @staticmethod
    def _table_matrix_key(table) -> str | None:
        from docx.oxml.ns import qn

        description = table._tbl.tblPr.find(qn("w:tblDescription"))
        return description.get(qn("w:val")) if description is not None else None

    @staticmethod
    def _table_case_id(table) -> str:
        match = re.match(r"(CP\d+)", table.rows[0].cells[0].text.strip())
        return match.group(1) if match else "CP01"

    def _build_case_id(self, document=None) -> str:
        self.report_folder.mkdir(parents=True, exist_ok=True)
        numbers = []
        for report in self.report_folder.glob("CP*.docx"):
            match = re.match(r"CP(\d+)", report.stem)
            if match:
                numbers.append(int(match.group(1)))

        if document is not None:
            for paragraph in document.paragraphs:
                numbers.extend(
                    int(match.group(1))
                    for match in re.finditer(r"\bCP(\d+)\b", paragraph.text)
                )

            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        numbers.extend(
                            int(match.group(1))
                            for match in re.finditer(r"\bCP(\d+)\b", cell.text)
                        )

        next_number = max(numbers, default=0) + 1
        return f"CP{next_number:02d}"

    def _configure_document(self, document) -> None:
        section = document.sections[0]
        section.top_margin = self._Inches(0.6)
        section.bottom_margin = self._Inches(0.6)
        section.left_margin = self._Inches(0.65)
        section.right_margin = self._Inches(0.65)

        normal = document.styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = self._Pt(9)

    def _add_case_table(
        self,
        document,
        case_id: str,
        execution_folder: Path,
        qa_actions: list[dict],
        start_url: str | None,
        case_name: str | None = None,
        matrix_case_key: str | None = None,
        matrix_case_id: str | None = None,
        expected_result: str | None = None,
        obtained_result: str | None = None,
    ):
        host = urlparse(start_url or "").netloc or "sistema web"
        first_value = next(
            (
                action.get("value")
                for action in qa_actions
                if action.get("value") and action.get("value") != "[MASKED]"
            ),
            None
        )
        if case_name:
            name = case_name.strip()
        else:
            name = f"Validación del recorrido en {host}"
            if first_value:
                name += f" - {first_value}"
            if not qa_actions:
                name = f"Ejecución incompleta en {host}"

        separator = None
        if document.tables:
            separator = document.add_paragraph()
            separator.paragraph_format.space_after = self._Pt(8)

        table = document.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        table.alignment = self._WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if matrix_case_key:
            description = self._OxmlElement("w:tblDescription")
            description.set(self._qn("w:val"), matrix_case_key)
            table._tbl.tblPr.append(description)

        title = table.rows[0].cells[0]
        title.text = f"{case_id} - {name}"
        self._style_cell(title, "102A43", "FFFFFF", 14)

        self._add_field_row(
            table,
            "ID",
            case_id
        )
        self._add_field_row(
            table,
            "Nombre del caso de prueba",
            name
        )
        if matrix_case_id:
            self._add_field_row(table, "ID de matriz", matrix_case_id)
        self._add_field_row(
            table,
            "Descripción",
            (
                "Verificar el flujo funcional registrado y comprobar que el "
                "sistema permita completar las acciones sin errores."
                if qa_actions
                else "Documentar una ejecución interrumpida antes de registrar acciones."
            )
        )

        preconditions = [
            "El sistema bajo prueba se encuentra operativo.",
            "El usuario cuenta con permisos para ejecutar el flujo.",
            "Los datos requeridos por la prueba se encuentran disponibles.",
        ]
        if start_url:
            preconditions.insert(0, f"El sistema está disponible en: {start_url}")
        self._add_list_row(table, "Precondiciones", preconditions)

        steps_cell = self._add_empty_field_row(table, "Pasos a seguir")
        if not qa_actions:
            empty_paragraph = steps_cell.add_paragraph(
                "No se registraron pasos antes de cerrar el navegador."
            )
            empty_paragraph.italic = True

        for action in qa_actions:
            paragraph = steps_cell.add_paragraph()
            paragraph.paragraph_format.space_after = self._Pt(4)
            run = paragraph.add_run(
                f"{action.get('sequence', '')}. "
                f"{action.get('description', 'Ejecutar acción')}"
            )
            run.bold = True

            evidence = self._evidence_path(execution_folder, action)
            if evidence and evidence.exists():
                evidence_paragraph = steps_cell.add_paragraph()
                evidence_paragraph.add_run().add_picture(
                    str(evidence),
                    width=self._Inches(5.8)
                )

        if expected_result and expected_result.strip():
            self._add_field_row(
                table,
                "Resultado esperado",
                expected_result.strip(),
            )
        elif matrix_case_key:
            self._add_field_row(
                table,
                "Resultado esperado",
                "No se tiene en la matriz.",
            )
        else:
            expected = [
                "El flujo se completa correctamente.",
                "No se muestran mensajes de error.",
                "Cada acción produce el comportamiento esperado del sistema.",
            ]
            if not qa_actions:
                expected = [
                    "La ejecución debe iniciar y registrar al menos un paso.",
                    "Si el navegador se cierra, la ejecución debe quedar identificada como incompleta.",
                ]
            self._add_list_row(table, "Resultado esperado", expected)
        default_obtained = (
            "El recorrido fue registrado con sus evidencias. La validación "
            "automática del resultado funcional queda pendiente de la reproducción del flujo."
            if qa_actions
            else "El navegador se cerró antes de registrar acciones; no fue posible evaluar el recorrido."
        )
        self._add_field_row(
            table,
            "Resultado obtenido",
            obtained_result.strip()
            if obtained_result and obtained_result.strip()
            else default_obtained,
        )
        self._add_field_row(
            table,
            "Estado",
            "OBSERVABLE" if qa_actions else "INCOMPLETO"
        )
        return table, separator

    def _add_field_row(self, table, label: str, value: str) -> None:
        row = table.add_row()
        cell = row.cells[0]
        cell.text = ""
        self._shade_cell(cell, "FFFFFF", "243B53")
        paragraph = cell.paragraphs[0]
        label_run = paragraph.add_run(f"{label}:\n")
        label_run.bold = True
        label_run.font.size = self._Pt(9)
        label_run.font.color.rgb = self._RGBColor(36, 59, 83)
        value_run = paragraph.add_run(value)
        value_run.font.size = self._Pt(9)
        value_run.font.color.rgb = self._RGBColor(36, 59, 83)

    def _add_empty_field_row(self, table, label: str):
        row = table.add_row()
        cell = row.cells[0]
        cell.text = ""
        self._shade_cell(cell, "FFFFFF", "243B53")
        paragraph = cell.paragraphs[0]
        label_run = paragraph.add_run(f"{label}:\n")
        label_run.bold = True
        label_run.font.size = self._Pt(9)
        label_run.font.color.rgb = self._RGBColor(36, 59, 83)
        return cell

    def _add_list_row(self, table, label: str, values: list[str]) -> None:
        cell = self._add_empty_field_row(table, label)
        for index, value in enumerate(values):
            paragraph = cell.add_paragraph()
            paragraph.style = "List Bullet"
            paragraph.add_run(value)

    def _style_cell(
        self,
        cell,
        fill: str,
        color: str,
        size: int,
        text: str | None = None
    ) -> None:
        if text is not None:
            cell.text = text
        self._shade_cell(cell, fill, color)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = self._Pt(size)
                run.font.color.rgb = self._RGBColor.from_string(color)

    @staticmethod
    def _evidence_path(execution_folder: Path, action: dict) -> Path | None:
        screenshot = action.get("screenshot")
        if not screenshot:
            return None
        return execution_folder / Path(screenshot.replace("\\", "/"))

    def _shade_cell(self, cell, fill: str, color: str | None = None) -> None:
        properties = cell._tc.get_or_add_tcPr()
        shading = self._OxmlElement("w:shd")
        shading.set(self._qn("w:fill"), fill)
        properties.append(shading)

        if color:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = self._RGBColor.from_string(color)
